#!/usr/bin/env python3
"""Post-process pandoc's thesis.docx into the final layout (thesis/PLAN.md §1.2–1.4, §4.2).

  * title page from docs/formatting_examples/Appendix_A_-_Sample_Title_Page_Formatting.docx:
    its paragraph styles (Т1…Т11) are copied and filled from thesis/metadata.yaml;
  * ЗМІСТ + TOC field before the first level-1 heading that is not АНОТАЦІЯ / ABSTRACT
    (Word fills it: Ctrl+A, F9 — the document also asks Word to update fields on open);
  * page number top-right in the header, the title page counted but not numbered
    («different first page»);
  * every chapter / structural element on a new page (Heading 1 «page break before»);
  * tables ruled, full text width with content-fitted columns, header row centred (§1.7);
  * «ДОДАТОК А» / title of the appendix on two lines (§1.4);
  * the 24 pt gap after tables, listings, figure captions, formulas, scenarios and headings as
    `before` on the next paragraph (§0.5, «max, not sum»).

Usage: assemble.py IN.docx OUT.docx [--metadata thesis/metadata.yaml] [--template T.docx]
"""
import argparse
import copy
import re
import sys
from pathlib import Path

import yaml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

THESIS = Path(__file__).resolve().parent.parent
REPO = THESIS.parent
TEMPLATE = REPO / "docs/formatting_examples/Appendix_A_-_Sample_Title_Page_Formatting.docx"
FRONT_MATTER = {"анотація", "abstract"}  # headings that come before ЗМІСТ (PLAN.md §1.3)


def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), v)
    return e


# --------------------------------------------------------------------------- title page
def title_lines(meta):
    """(template style name, text) per paragraph, mirroring the Appendix_A template."""
    m = {k: ("" if v is None else str(v)) for k, v in meta.items() if not isinstance(v, list)}
    lines = [("Т1", m["ministry"]), ("Т2", m["university"]), ("Т3", m["institute"]),
             ("Т4", m["department"]), ("Т5", ""), ("Т5", ""), ("Т5", ""),
             ("Т6", m["student"]), ("Т7", f"(група {m['group']})"), ("Т5", ""),
             ("Т8", m["work_type"]), ("Т9", m["title_uk"]), ("Т5", ""),
             ("Т10", "Спеціальність:"), ("Т11", m["specialty"]), ("Т5", ""),
             ("Т10", "Освітньо-професійна програма:"), ("Т11", m["program"]), ("Т5", ""),
             ("Т10", "Керівник:"), ("Т11", m["supervisor"])]
    consultants = meta.get("consultants") or []
    if consultants:
        lines += [("Т5", ""), ("Т10", "Консультанти:")] + [("Т11", str(c)) for c in consultants]
    lines += [("Т5", ""), ("Т5", ""), ("Т5", ""), ("Т11", m["city_year"])]
    return lines


def copy_title_styles(doc, template):
    """Copy the template's Т1…Т11 styles under fresh ids; return {name: styleId}."""
    src = Document(str(template)).styles.element
    dst = doc.styles.element
    taken = {s.get(qn("w:styleId")) for s in dst.findall(qn("w:style"))}
    ids = {}
    for s in src.findall(qn("w:style")):
        name_el = s.find(qn("w:name"))
        name = name_el.get(qn("w:val")) if name_el is not None else ""
        if not (name.startswith("Т") and name[1:].isdigit()):
            continue
        new = copy.deepcopy(s)
        sid = f"TitlePage{name[1:]}"
        assert sid not in taken, sid
        new.set(qn("w:styleId"), sid)
        for tag in ("w:basedOn", "w:link", "w:rsid"):  # standalone: no template parents
            for e in new.findall(qn(tag)):
                new.remove(e)
        ppr = new.find(qn("w:pPr"))
        if ppr is None:
            ppr = el("w:pPr")
            new.append(ppr)
        for tag in ("w:pageBreakBefore", "w:jc", "w:ind"):
            for e in ppr.findall(qn(tag)):
                ppr.remove(e)
        # The template centres everything through docDefaults; make that explicit here.
        ppr.append(el("w:ind", **{"w:left": "0", "w:right": "0", "w:firstLine": "0"}))
        ppr.append(el("w:jc", **{"w:val": "center"}))
        # The template has no space before/after (its docDefaults); reference.docx's docDefaults
        # add 10 pt after, which pushed «Одеса – рік» onto page 2.
        spacing = ppr.find(qn("w:spacing"))
        if spacing is None:
            spacing = el("w:spacing")
            ppr.find(qn("w:ind")).addprevious(spacing)  # schema order: spacing, ind, jc
        for side in ("w:before", "w:after"):
            if spacing.get(qn(side)) is None:
                spacing.set(qn(side), "0")
        dst.append(new)
        ids[name] = sid
    missing = {f"Т{i}" for i in range(1, 12)} - ids.keys()
    if missing:
        sys.exit(f"assemble.py: template lacks styles {sorted(missing)}")
    return ids


def make_para(style_id, text):
    p = el("w:p")
    ppr = el("w:pPr")
    ppr.append(el("w:pStyle", **{"w:val": style_id}))
    p.append(ppr)
    if text:
        r = el("w:r")
        t = el("w:t")
        t.text = text
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        r.append(t)
        p.append(r)
    return p


def insert_title_page(doc, meta, template):
    ids = copy_title_styles(doc, template)
    body = doc.element.body
    first = body[0]
    for style_name, text in title_lines(meta):
        first.addprevious(make_para(ids[style_name], text))
    # What follows the title page starts on page 2 even if it is not a Heading 1.
    if first.tag == qn("w:p"):
        ppr = first.get_or_add_pPr()
        if ppr.find(qn("w:pageBreakBefore")) is None:
            ppr.insert(1 if ppr.find(qn("w:pStyle")) is not None else 0, el("w:pageBreakBefore"))
    else:
        brk = el("w:p")
        r = el("w:r")
        r.append(el("w:br", **{"w:type": "page"}))
        brk.append(r)
        first.addprevious(brk)


# --------------------------------------------------------------------------- fields
def field_runs(instr, placeholder=""):
    runs = []
    r = el("w:r")
    r.append(el("w:fldChar", **{"w:fldCharType": "begin", "w:dirty": "true"}))
    runs.append(r)
    r = el("w:r")
    it = el("w:instrText")
    it.text = f" {instr} "
    it.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    r.append(it)
    runs.append(r)
    r = el("w:r")
    r.append(el("w:fldChar", **{"w:fldCharType": "separate"}))
    runs.append(r)
    r = el("w:r")
    t = el("w:t")
    t.text = placeholder
    r.append(t)
    runs.append(r)
    r = el("w:r")
    r.append(el("w:fldChar", **{"w:fldCharType": "end"}))
    runs.append(r)
    return runs


def is_heading1(p):
    style = p.find(qn("w:pPr") + "/" + qn("w:pStyle"))
    return style is not None and style.get(qn("w:val")) == "Heading1"


def para_text(p):
    return "".join(t.text or "" for t in p.iter(qn("w:t"))).strip()


def insert_toc(doc):
    target = None
    for p in doc.element.body.iter(qn("w:p")):
        if is_heading1(p) and para_text(p).lower() not in FRONT_MATTER:
            target = p
            break
    if target is None:
        print("assemble.py: no chapter heading found, TOC not inserted", file=sys.stderr)
        return
    # АНОТАЦІЯ / ABSTRACT precede ЗМІСТ and are not listed in it (addendum_v): same look as
    # Heading 1 via the TOC Title style, but no outline level, so the TOC field skips them.
    for p in list(doc.element.body.iter(qn("w:p"))):
        if p is target:
            break
        if is_heading1(p) and para_text(p).lower() in FRONT_MATTER:
            p.find(qn("w:pPr") + "/" + qn("w:pStyle")).set(qn("w:val"), "TOCTitle")
    title = make_para("TOCTitle", "ЗМІСТ")
    toc = make_para("toc1", "")
    for r in field_runs('TOC \\o "1-2" \\h \\z \\u',
                        "Зміст буде сформовано після оновлення полів (Ctrl+A, F9)."):
        toc.append(r)
    target.addprevious(title)
    target.addprevious(toc)
    # Ask Word to refresh fields (TOC, PAGE) when the document is opened.
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        upd = el("w:updateFields", **{"w:val": "true"})
        # CT_Settings is an ordered sequence: insert before the first element that follows it.
        later = {"hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids",
                 "mathPr", "attachedSchema", "themeFontLang", "clrSchemeMapping",
                 "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures", "forceUpgrade",
                 "captions", "readModeInkLockDown", "smartTagType", "schemaLibrary",
                 "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator"}
        nxt = next((c for c in settings if c.tag.split("}")[-1] in later), None)
        if nxt is not None:
            nxt.addprevious(upd)
        else:
            settings.append(upd)


def number_pages(doc):
    for section in doc.sections:
        section.different_first_page_header_footer = True  # title page: counted, not printed
        section.first_page_header.is_linked_to_previous = False
        for p in section.first_page_header.paragraphs:
            for r in list(p.runs):
                r._r.getparent().remove(r._r)
        header = section.header
        header.is_linked_to_previous = False
        p = header.paragraphs[0]
        p.style = doc.styles["Header"]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.first_line_indent = Mm(0)
        for r in field_runs("PAGE", "1"):
            rpr = el("w:rPr")
            rpr.append(el("w:rFonts", **{"w:ascii": "Times New Roman", "w:hAnsi": "Times New Roman",
                                         "w:cs": "Times New Roman"}))
            rpr.append(el("w:sz", **{"w:val": str(int(Pt(14).pt * 2))}))
            r.insert(0, rpr)
            p._p.append(r)


# --------------------------------------------------------------------------- tables
TEXT_WIDTH = 9921          # twips: 210 − 25 − 10 = 175 mm
CELL_PAD = 2 * 115         # default left + right cell margins (twips)
CHAR_W = {28: 150, 24: 130}  # glyph width (twips) of TNR 14 / 12 pt, Cyrillic, with margin
WIDE_COLS = 5              # PLAN.md §1.7: 12 pt allowed for wide tables


def word_glyphs(word):
    """Width of a word in average glyphs: capitals (Ф, В, Ш…) are much wider than the average."""
    return sum(1.35 if ch.isupper() else 1 for ch in word)


def column_widths(cols, size):
    """Auto-fit like a browser: each column gets at least its longest word, the rest of the text
    width is shared in proportion to how much longer the column's text is than that minimum."""
    cw = CHAR_W[size]
    # ≥ 3 average glyphs per word: short IDs with a wide Cyrillic letter («П1», «Д10») must not wrap;
    # capitals count 1.35 glyphs, else «ФВ-12» breaks after its hyphen in a column fitted to it
    mins = [max((max(word_glyphs(w), 3) for c in col for w in c.split()), default=3) * cw + CELL_PAD
            for col in cols]
    maxs = [max(max((len(c) for c in col), default=1) * cw + CELL_PAD, m) for col, m in zip(cols, mins)]
    if sum(maxs) <= TEXT_WIDTH:
        return [round(m * TEXT_WIDTH / sum(maxs)) for m in maxs], True
    if sum(mins) >= TEXT_WIDTH:
        return [round(m * TEXT_WIDTH / sum(mins)) for m in mins], False
    spare, extra = TEXT_WIDTH - sum(mins), [b - a for a, b in zip(mins, maxs)]
    return [round(a + spare * e / sum(extra)) for a, e in zip(mins, extra)], True


def format_tables(doc):
    """PLAN.md §1.7: ruled tables over the full text width, header row centred and repeated on
    every page, single spacing (Compact style); 12 pt in wide tables (≥ 5 columns or words that
    do not fit at 14 pt)."""
    for tbl in doc.element.body.iter(qn("w:tbl")):
        rows = tbl.findall(qn("w:tr"))
        ncols = len(tbl.find(qn("w:tblGrid")).findall(qn("w:gridCol")))
        cols = [[] for _ in range(ncols)]
        for tr in rows:
            for k, tc in enumerate(tr.findall(qn("w:tc"))[:ncols]):
                cols[k].append(" ".join(para_text(p) for p in tc.iter(qn("w:p"))))
        size = 24 if ncols >= WIDE_COLS else 28
        widths, fits = column_widths(cols, size)
        if not fits and size == 28:
            size = 24
            widths, _ = column_widths(cols, size)
        widths[-1] += TEXT_WIDTH - sum(widths)

        tblpr = tbl.find(qn("w:tblPr"))
        for tag in ("w:tblW", "w:tblBorders", "w:tblLayout"):
            for e in tblpr.findall(qn(tag)):
                tblpr.remove(e)
        style = tblpr.find(qn("w:tblStyle"))
        pos = list(tblpr).index(style) + 1 if style is not None else 0
        borders = el("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            borders.append(el(f"w:{side}", **{"w:val": "single", "w:sz": "4", "w:space": "0",
                                              "w:color": "000000"}))
        for e in reversed([el("w:tblW", **{"w:w": str(TEXT_WIDTH), "w:type": "dxa"}),
                           el("w:jc", **{"w:val": "center"}), borders,
                           el("w:tblLayout", **{"w:type": "fixed"})]):
            for old in tblpr.findall(e.tag):
                tblpr.remove(old)
            tblpr.insert(pos, e)
        grid = tbl.find(qn("w:tblGrid"))
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(w))
        for i, tr in enumerate(rows):
            for k, tc in enumerate(tr.findall(qn("w:tc"))):
                tcpr = tc.find(qn("w:tcPr"))
                if tcpr is None:
                    tcpr = el("w:tcPr")
                    tc.insert(0, tcpr)
                for e in tcpr.findall(qn("w:tcW")):
                    tcpr.remove(e)
                if k < len(widths):
                    tcpr.insert(0, el("w:tcW", **{"w:w": str(widths[k]), "w:type": "dxa"}))
                for p in tc.iter(qn("w:p")):
                    if i == 0:
                        ppr = p.find(qn("w:pPr"))
                        if ppr is None:
                            ppr = el("w:pPr")
                            p.insert(0, ppr)
                        for e in ppr.findall(qn("w:jc")):
                            ppr.remove(e)
                        ppr.append(el("w:jc", **{"w:val": "center"}))
                    if size != 28:
                        for r in p.iter(qn("w:r")):
                            rpr = r.find(qn("w:rPr"))
                            if rpr is None:
                                rpr = el("w:rPr")
                                r.insert(0, rpr)
                            for tag in ("w:sz", "w:szCs"):
                                for e in rpr.findall(qn(tag)):
                                    rpr.remove(e)
                            rpr.append(el("w:sz", **{"w:val": str(size)}))
                            rpr.append(el("w:szCs", **{"w:val": str(size)}))


# --------------------------------------------------------------------------- vertical spacing
# PLAN.md §0.5: the gap *after* an element, in twentieths of a point (24 pt = 480). Styles keep
# `after` = 0 (make_reference_docx.py); the gap becomes `before` on the next paragraph, raised
# only if that paragraph's own `before` is smaller — «max, not sum».
GAP_AFTER = {
    "tbl": 480,             # table (no paragraph of its own)
    "SourceCode": 480,      # listing: last code line
    "ImageCaption": 480,    # figure caption
    "Formula": 480,         # formula line, also before «де …»
    "ScenarioLast": 480,    # closing rule of a scenario
    "ScenarioSingle": 480,
    "Heading2": 480,        # subsection heading: one empty line after
    "Heading3": 480,
    "Heading1": 420,        # chapter heading: 21 pt, as before
}


def style_spacing(doc):
    """styleId -> effective `before` (twips) of every paragraph style, following basedOn."""
    styles = {s.get(qn("w:styleId")): s for s in doc.styles.element.findall(qn("w:style"))}
    default = doc.styles.element.find(f"{qn('w:docDefaults')}/{qn('w:pPrDefault')}/"
                                      f"{qn('w:pPr')}/{qn('w:spacing')}")
    base = int(default.get(qn("w:before"), "0")) if default is not None else 0
    cache = {}

    def before(sid, depth=0):
        if sid in cache:
            return cache[sid]
        st = styles.get(sid)
        if st is None or depth > 20:
            return base
        sp = st.find(f"{qn('w:pPr')}/{qn('w:spacing')}")
        if sp is not None and sp.get(qn("w:before")) is not None:
            val = int(sp.get(qn("w:before")))
        else:
            parent = st.find(qn("w:basedOn"))
            val = before(parent.get(qn("w:val")), depth + 1) if parent is not None else base
        cache[sid] = val
        return val

    normal = next((sid for sid, st in styles.items() if st.get(qn("w:default")) == "1"
                   and st.get(qn("w:type")) == "paragraph"), "Normal")
    return lambda sid: before(sid or normal)


def p_style(p):
    s = p.find(qn("w:pPr") + "/" + qn("w:pStyle"))
    return s.get(qn("w:val")) if s is not None else ""


def set_before(p, twips):
    ppr = p.get_or_add_pPr()
    sp = ppr.find(qn("w:spacing"))
    if sp is None:
        sp = el("w:spacing")
        # CT_PPrBase: spacing comes after pStyle…tabs…, before ind/jc/rPr; find the first later one.
        later = {qn(t) for t in ("w:ind", "w:contextualSpacing", "w:mirrorIndents",
                                 "w:suppressOverlap", "w:jc", "w:textDirection",
                                 "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
                                 "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")}
        nxt = next((c for c in ppr if c.tag in later), None)
        if nxt is not None:
            nxt.addprevious(sp)
        else:
            ppr.append(sp)
    sp.set(qn("w:before"), str(twips))


def space_blocks(doc):
    """Write the §0.5 gap after tables, listings, figures, formulas, scenarios and headings as
    `before` on the following body paragraph (max of the two, never their sum)."""
    before_of = style_spacing(doc)
    blocks = [e for e in doc.element.body if e.tag in (qn("w:p"), qn("w:tbl"))]
    for prev, cur in zip(blocks, blocks[1:]):
        key = "tbl" if prev.tag == qn("w:tbl") else p_style(prev)
        gap = GAP_AFTER.get(key)
        if not gap or cur.tag != qn("w:p") or p_style(cur) == "Heading1":
            continue
        direct = cur.find(f"{qn('w:pPr')}/{qn('w:spacing')}")
        own = (int(direct.get(qn("w:before"))) if direct is not None
               and direct.get(qn("w:before")) is not None else before_of(p_style(cur)))
        if own < gap:
            set_before(cur, gap)


def split_appendix_headings(doc):
    """PLAN.md §1.4 and both examples: «ДОДАТОК А» on one line, the title below it in sentence
    case (Heading 1 is all caps), as one heading so ЗМІСТ lists «Додаток А Лістинг програми»."""
    for p in doc.element.body.iter(qn("w:p")):
        if not is_heading1(p):
            continue
        m = re.match(r"(Додаток\s+\S)\s+(.+)$", para_text(p))
        if not m:
            continue
        for r in p.findall(qn("w:r")):
            p.remove(r)
        r = el("w:r")
        t = el("w:t")
        t.text = m.group(1)
        r.append(t)
        r.append(el("w:br"))
        p.append(r)
        r = el("w:r")
        rpr = el("w:rPr")
        rpr.append(el("w:caps", **{"w:val": "0"}))
        r.append(rpr)
        t = el("w:t")
        t.text = m.group(2)
        r.append(t)
        p.append(r)


def chapters_on_new_page(doc):
    h1 = doc.styles["Heading 1"]
    h1.paragraph_format.page_break_before = True


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("input")
    ap.add_argument("output")
    ap.add_argument("--metadata", default=str(THESIS / "metadata.yaml"))
    ap.add_argument("--template", default=str(TEMPLATE))
    a = ap.parse_args()

    meta = yaml.safe_load(Path(a.metadata).read_text(encoding="utf-8"))
    doc = Document(a.input)
    insert_title_page(doc, meta, a.template)
    insert_toc(doc)
    number_pages(doc)
    format_tables(doc)
    split_appendix_headings(doc)
    chapters_on_new_page(doc)
    space_blocks(doc)
    doc.save(a.output)


if __name__ == "__main__":
    main()
